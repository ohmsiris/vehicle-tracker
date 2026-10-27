"""
The stock list as an Excel sheet to fill in the minimum quantity (and the shelf) for every item, and to add new items below.

    .venv/Scripts/python.exe scripts/export_stock_sheet.py "<output .xlsx>"

Yellow cells are for the owner. The sheet comes back through scripts/import_stock_sheet.py (rows with a รหัส update the minimum / shelf; rows without one are
new items). Column A (รหัส) is the item's id: do not change it.
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook                                       # noqa: E402
from openpyxl.comments import Comment                                # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side   # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation         # noqa: E402

from app import stock                                                 # noqa: E402

FONT = "Arial"
HEAD_ROW = 7
COLS = [("รหัส (ห้ามแก้)", 12), ("หมวด", 14), ("ชนิด", 22), ("ยี่ห้อ", 14), ("รุ่น", 16), ("สเปค", 22), ("หน่วย", 9), ("คงเหลือตอนนี้", 14),
        ("ขั้นต่ำ", 11), ("ที่เก็บ / ชั้น", 16), ("หมายเหตุ", 34)]
YELLOW = PatternFill("solid", fgColor="FFF2A8")
GREY = PatternFill("solid", fgColor="EDEDED")
HEAD = PatternFill("solid", fgColor="1F3A5F")
THIN = Side(style="thin", color="BBBBBB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def main(out: str) -> None:
    ov = stock.overview()
    items = sorted((i for i in ov["items"] if i.get("active") is not False),
                   key=lambda i: tuple(str(i.get(k) or "") for k in ("category", "kind", "brand", "model", "spec", "name")))
    wb = Workbook()
    ws = wb.active
    ws.title = "ขั้นต่ำสต็อก"
    f = lambda **kw: Font(name=FONT, size=kw.pop("size", 10), **kw)
    ws["A1"] = "รายการสต็อกห้องอะไหล่ช่าง — กรอกขั้นต่ำ"
    ws["A1"].font = f(size=14, bold=True)
    legend = [
        "ช่องสีเหลือง = ให้คุณกรอก: “ขั้นต่ำ” (ต่ำกว่านี้ระบบขึ้นว่าควรสั่งเพิ่ม) และ “ที่เก็บ / ชั้น”. เว้นว่างไว้ = ไม่เปลี่ยน.",
        "ของใหม่: เพิ่มแถวต่อท้ายตาราง ปล่อย “รหัส” ว่าง กรอก หมวด / ชนิด / ยี่ห้อ / รุ่น / สเปค / หน่วย — “คงเหลือตอนนี้” ของแถวใหม่คือยอดตั้งต้น. ชื่อมาตรฐานระบบสร้างให้เอง (ชนิด + ยี่ห้อ + รุ่น + สเปค).",
        "ของที่มีอยู่แล้ว: คอลัมน์อื่นนอกจากสีเหลืองแก้ในนี้ไม่มีผล (จำนวนในสต็อกนับด้วยปุ่ม “นับ” ในหน้าสต็อก). แถวตัวอย่างสีเทาด้านล่างไม่ถูกนำเข้า.",
    ]
    for n, text in enumerate(legend, start=2):
        ws.cell(row=n, column=1, value=text).font = f(italic=True, color="444444")
    ws.cell(row=5, column=1, value=f"{len(items)} รายการ ณ วันที่ส่งออก").font = f(color="444444")
    for c, (name, width) in enumerate(COLS, start=1):
        cell = ws.cell(row=HEAD_ROW, column=c, value=name)
        cell.font = f(bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX
        ws.column_dimensions[ws.cell(row=HEAD_ROW, column=c).column_letter].width = width
    ws.row_dimensions[HEAD_ROW].height = 30
    r = HEAD_ROW + 1
    for i in items:
        row = [i["id"], i.get("category"), i.get("kind") or i["name"], i.get("brand"), i.get("model"), i.get("spec"), i.get("unit"), i["on_hand"],
               float(i["min_qty"]) if i.get("min_qty") is not None else None, i.get("location"), i.get("notes")]
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = f(color="888888" if c == 1 else "000000", size=8 if c == 1 else 10)
            cell.border = BOX
            if c in (9, 10):
                cell.fill = YELLOW
            elif c == 8:
                cell.number_format = "#,##0.##"
        r += 1
    last = r - 1
    # a few empty yellow rows for new items, and one example row (not imported)
    for _ in range(10):
        for c in range(1, len(COLS) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BOX
            cell.font = f()
            if c in (2, 3, 4, 5, 6, 7, 8, 9, 10):
                cell.fill = YELLOW
        r += 1
    example = [None, "ไฟฟ้า", "แมกเนติก", "Fuji", "SC-N3", "220V", "ตัว", 2, 1, "ตู้ A", "ตัวอย่าง — ไม่ถูกนำเข้า (ของใหม่: เว้นรหัสว่าง)"]
    for c, v in enumerate(example, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = f(italic=True, color="777777")
        cell.fill = GREY
        cell.border = BOX
    ws.cell(row=HEAD_ROW, column=9).comment = Comment("ต่ำกว่าเลขนี้ → หน้าสต็อกขึ้น “ควรสั่งเพิ่ม”. เว้นว่าง = ไม่ตั้งขั้นต่ำ", "ระบบ")
    dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True, showErrorMessage=True, errorTitle="ขั้นต่ำ", error="ใส่ตัวเลขตั้งแต่ 0 ขึ้นไป")
    ws.add_data_validation(dv)
    dv.add(f"I{HEAD_ROW + 1}:I{r - 1}")
    ws.freeze_panes = ws.cell(row=HEAD_ROW + 1, column=3)
    ws.auto_filter.ref = f"A{HEAD_ROW}:{ws.cell(row=HEAD_ROW, column=len(COLS)).column_letter}{last}"
    ws.sheet_view.zoomScale = 110
    wb.save(out)
    print(f"{len(items)} items -> {out}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
