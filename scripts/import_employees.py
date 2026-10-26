"""
One-time (re-runnable) import of the staff into the employees table (schema_v21.sql).

    .venv/Scripts/python.exe scripts/import_employees.py --roster Roster.csv --kk "เงินเดือนแก่งคอย ....xlsx" --drivers "เงินเดือน พขร. ....xlsx"          (dry run: prints the plan)
    ... the same with --apply                                                                                                                           (writes)

  --roster   the Employee Roster sheet exported as CSV (ลำดับ, สาขา, รหัสพนักงาน, ชื่อ-สกุล, ตำแหน่ง/แผนก): the main list, Saraburi. Authoritative.
  --kk       the Kaeng Khoi pay workbook (sheet "สรุป"): code, name, department of the Kaeng Khoi staff -- the roster has none of them.
  --drivers  the drivers' pay workbook (sheet "ตารางรายรับ-จ่ายพนง."): the start date of each SB-PD driver (written in พ.ศ. there).

A code that is already in the table is left alone (so it can be run again and again). A พขร. is linked to the `drivers` row with the same
name when exactly one matches. Nothing is read from or written to git: the files stay where they are.
"""
import argparse
import csv
import re
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app import employees, fleet, supabase                      # noqa: E402
from app.employees import split_prefix                          # noqa: E402

POSITION_ALIASES = {"เสมียน": "เสมียนออฟฟิศ", "พขร.": "พนักงานขับรถ", "พขร.รถสาย": "พนักงานขับรถสาย", "พขร.รถหน้าโรง": "พนักงานขับรถหน้าโรง"}


def read_roster(path: str) -> list[dict]:
    out = []
    with open(path, encoding="utf-8-sig", newline="") as f:
        for row in csv.reader(f):
            if len(row) < 5 or not re.match(r"^SB-", row[2].strip()):
                continue
            branch = row[1].replace("สาขา", "").strip() or "สระบุรี"
            position = POSITION_ALIASES.get(row[4].strip(), row[4].strip())
            e = {"code": row[2].strip().upper(), "full_name": row[3].strip(), "branch": branch, "position": position or None, "source": "roster"}
            if position == "พ้นสภาพ":                           # the roster marks someone who has left like this
                e.update(position=None, resigned=True, notes="พ้นสภาพ ตามรายชื่อพนักงาน")
            out.append(e)
    return out


def read_kk(path: str) -> list[dict]:
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True)["สรุป"]
    out = []
    for r in range(5, ws.max_row + 1):
        v = [c.value for c in ws[r]]
        code, name, dept = v[1], v[2], v[3]
        if not (code and str(code).startswith("SB-KK") and name and str(name).strip()):
            continue                                            # a code with no name is an empty seat
        name, dept = re.sub(r"\s+", " ", str(name)).strip(), (str(dept).strip() if dept else None)
        e = {"code": str(code).strip().upper(), "full_name": name, "branch": "แก่งคอย", "position": POSITION_ALIASES.get(dept, dept), "source": "kk"}
        if re.search(r"-\s*ออก\s*$", name):                     # "นายอาจอ-ออก": written on the sheet by whoever noticed
            e.update(full_name=re.sub(r"\s*-\s*ออก\s*$", "", name), resigned=True, notes="มีคำว่า “-ออก” ต่อท้ายชื่อในใบเงินเดือน")
        out.append(e)
    return out


def read_start_dates(path: str) -> dict:
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True)["ตารางรายรับ-จ่ายพนง."]
    out = {}
    for r in range(5, ws.max_row + 1):
        v = [c.value for c in ws[r]]
        if len(v) > 36 and v[1] and str(v[1]).startswith("SB-") and v[34] and v[35] and v[36]:
            try:
                out[str(v[1]).strip().upper()] = date(int(v[36]) - 543, int(v[35]), int(v[34])).isoformat()
            except (ValueError, TypeError):
                pass
    return out


def fix_prefixes(apply: bool) -> None:
    """People already in the table with no prefix and no gender: read the title (and the gender) out of the name. Never touches one that has a prefix."""
    rows = [r for r in supabase.select_all("employees", {"select": "id,code,full_name,prefix,gender", "order": "code.asc"}) if not r.get("prefix") and not r.get("gender")]
    changes, unknown = [], []
    for r in rows:
        prefix, name, gender = split_prefix(r["full_name"])
        if prefix or gender:
            changes.append((r, {"prefix": prefix, "full_name": name, "gender": gender}))
        else:
            unknown.append((r["code"], r["full_name"]))
    print(f"to read: {len(rows)} · detected: {len(changes)} · not detectable: {len(unknown)}")
    for code, name in unknown:
        print("   no title / gender in the name:", code, name)
    if apply:
        for r, values in changes:
            supabase.update("employees", {"id": f"eq.{r['id']}"}, values)
        print(f"updated: {len(changes)}")
    else:
        print("dry run -- nothing written (add --apply)")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--roster", required=True)
    ap.add_argument("--kk")
    ap.add_argument("--drivers")
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--fix-prefixes", action="store_true", help="only read the title / gender of people already in the table whose prefix is empty")
    args = ap.parse_args()
    if args.fix_prefixes:
        return fix_prefixes(args.apply)

    people = read_roster(args.roster)
    if args.kk:
        have = {p["code"] for p in people}
        people += [p for p in read_kk(args.kk) if p["code"] not in have]
    starts = read_start_dates(args.drivers) if args.drivers else {}
    for p in people:
        if p["code"] in starts:
            p["start_date"] = starts[p["code"]]

    existing = {e["code"].upper(): e for e in supabase.select_all("employees", {"select": "*", "order": "code.asc"})}
    drivers = supabase.select_all("drivers", {"select": "id,name", "order": "id.asc"})
    by_name: dict = {}
    for d in drivers:
        by_name.setdefault(fleet._name_key(d["name"]), []).append(d["id"])

    todo, linked, skipped = [], 0, 0
    for p in people:
        if p["code"] in existing:
            skipped += 1
            continue
        prefix, name, gender = split_prefix(p["full_name"])         # "นายจรูญ มินฉิมพลี" -> นาย + "จรูญ มินฉิมพลี" (+ the gender)
        row = {"code": p["code"], "prefix": prefix, "gender": gender, "full_name": name, "branch": p["branch"], "position": p.get("position"),
               "start_date": p.get("start_date"), "notes": p.get("notes"), "status": "resigned" if p.get("resigned") else "active", "driver_id": None}
        if employees.is_driver_position(p.get("position")):
            ids = by_name.get(fleet._name_key(p["full_name"]), [])
            if len(ids) == 1:
                row["driver_id"] = ids[0]
                linked += 1
        todo.append(row)

    by_branch: dict = {}
    for r in todo:
        by_branch[r["branch"]] = by_branch.get(r["branch"], 0) + 1
    print(f"people in the files: {len(people)}  ·  already in the table: {skipped}  ·  to add: {len(todo)} {by_branch}  ·  linked to a driver: {linked}  ·  with a start date: {sum(1 for r in todo if r.get('start_date'))}")
    print(f"marked as resigned: {[(r['code'], r['full_name']) for r in todo if r.get('status') == 'resigned']}")
    drv_unlinked = [r["code"] for r in todo if employees.is_driver_position(r.get("position")) and not r.get("driver_id")]
    print(f"drivers with no single matching driver row ({len(drv_unlinked)}): {drv_unlinked}  (scripts/merge_drivers.py ties them up)")
    if not args.apply:
        print("dry run -- nothing written (add --apply)")
        return
    for i in range(0, len(todo), 100):
        supabase.insert("employees", todo[i:i + 100])
    print(f"written: {len(todo)}")


if __name__ == "__main__":
    main()
