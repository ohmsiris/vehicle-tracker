"""
Imports the machinery spreadsheet (Google Sheets "Saraburi Maintenence Sheet", downloaded as .xlsx) into the
database: the seven machines, the parts / specs index, and the equipment-change log.

    .venv/Scripts/python.exe scripts/import_machinery.py "Saraburi Maintenance Sheet.xlsx"            # dry run
    .venv/Scripts/python.exe scripts/import_machinery.py "Saraburi Maintenance Sheet.xlsx" --write    # does it

Needs schema_v9.sql to have been run and `pip install openpyxl`. Safe to run again: rows that already exist are
skipped (see app/machineryimport.py). The .xlsx holds supplier prices and company details: it is git-ignored.
"""
import argparse
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from app import machinery, machineryimport as mi, supabase      # noqa: E402


def sheet_rows(wb, name):
    if name not in wb.sheetnames:
        sys.exit(f"ไม่พบแท็บ “{name}” ในไฟล์")
    return [list(r) for r in wb[name].iter_rows(values_only=True)]


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("file")
    ap.add_argument("--write", action="store_true", help="actually write (default is a dry run)")
    args = ap.parse_args()
    try:
        import openpyxl
    except ImportError:
        sys.exit("ต้องติดตั้ง openpyxl ก่อน:  .venv/Scripts/python.exe -m pip install openpyxl")
    if not supabase.configured():
        sys.exit("Supabase ยังไม่ได้ตั้งค่า (.env)")
    if not machinery.has_v9():
        sys.exit("ต้องรัน schema_v9.sql ใน Supabase ก่อน")

    wb = openpyxl.load_workbook(args.file, data_only=True)
    parts = mi.parse_main(sheet_rows(wb, "อะไหล่ Main"))
    motors = mi.parse_motors(sheet_rows(wb, "Motor Index Main"))
    prv = mi.parse_prv(sheet_rows(wb, "Pressure Relief Valve"))
    changes = mi.parse_changes(sheet_rows(wb, "Saraburi_General"), "สระบุรี") + mi.parse_changes(sheet_rows(wb, "Kaeng Khoi"), "แก่งคอย")

    existing_machines = supabase.select_all("machines", {"select": "*", "order": "id.asc"})
    have = {(m["branch"], m["kind"], m.get("number"), mi.squash(m.get("make"))) for m in existing_machines}
    to_create = [m for m in mi.SARABURI_MACHINES if (m["branch"], m["kind"], m["number"], mi.squash(m["make"])) not in have]
    print(f"ในไฟล์: อะไหล่ {len(parts)} แถว · มอเตอร์ {len(motors)} ตัว · Pressure Relief Valve {len(prv)} แถว · ประวัติเปลี่ยนอะไหล่ {len(changes)} รายการ")
    print(f"เครื่องที่จะเพิ่ม: {len(to_create)} (มีอยู่แล้ว {len(existing_machines)})")

    if args.write:
        for m in to_create:
            machinery.create_machine(m)
        existing_machines = supabase.select_all("machines", {"select": "*", "order": "id.asc"})
    # in a dry run the machines that WOULD exist stand in (with made-up ids) so linking can be counted
    machines = existing_machines if args.write else existing_machines + [{**m, "id": f"new-{i}"} for i, m in enumerate(to_create)]

    plan = mi.build_plan(parts=parts, motors=motors, prv=prv, changes=changes, machines=machines,
                         existing_parts=supabase.select_all("machine_parts", {"select": "location,equipment,part,spec", "order": "id.asc"}),
                         existing_events=supabase.select_all("machine_events", {"select": "branch,event_date,equipment,summary", "order": "id.asc"}))
    linked = sum(1 for p in plan["parts"] if "machine_id" in p)
    print(f"อะไหล่ที่จะเพิ่ม: {len(plan['parts'])} แถว (ผูกกับเครื่อง {linked}, ข้าม/มีอยู่แล้ว {plan['parts_skipped']})")
    print(f"พื้นที่ที่ไม่ใช่เครื่อง (ไม่ผูก): {', '.join(plan['unlinked_locations']) or '-'}")
    print(f"รายการเปลี่ยนอะไหล่ที่จะเพิ่ม: {len(plan['events'])} (มีอยู่แล้ว {plan['events_skipped']})")

    if not args.write:
        print("\n(dry run — ยังไม่ได้เขียนอะไร  ใส่ --write เพื่อทำจริง)")
        return
    for i in range(0, len(plan["parts"]), 100):
        supabase.insert("machine_parts", plan["parts"][i:i + 100])
    for e in plan["events"]:
        supabase.insert("machine_events", [e])
    print("\nเสร็จแล้ว")


if __name__ == "__main__":
    main()
