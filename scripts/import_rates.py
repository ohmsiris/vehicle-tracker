"""
Starting daily rates (employee_rates, schema_v25.sql) read from the monthly pay workbooks, with the owner's corrections applied.

    .venv/Scripts/python.exe scripts/import_rates.py "<drivers .xlsx>" "<staff .xlsx>" "<Kaeng Khoi .xlsx>" [--from 2026-08-26] [--apply]

Each row's pay for the days is written as a formula  =E5*390  (days x rate): the 390 is the rate. A person is skipped when they have no such formula
(a typed amount) unless OVERRIDES gives their rate, and a person already holding a rate dated --from is left alone, so it can be run again.
The rate starts on --from (default: the start of the September 2569 pay period, 26 Aug) as "ปกติ". Also applied here: the people paid by HQ,
a position correction, and the rates flagged for the owner to look at (REVIEW).
"""
import argparse
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app import employees, rates, supabase                      # noqa: E402

RATE_RE = re.compile(r"^=\s*\$?[A-Z]+\$?\d+\s*\*\s*([\d.]+)\s*$")
# the owner's corrections (Oct 2026)
OVERRIDES = {
    "SB-FR034": (390.0, "ค่าแรงปกติ 390 (HQ จ่าย 344 บางวัน: ส.ค. 1 วัน, ก.ย. 2 วัน)"),
    "SB-FR076": (390.0, "ค่าแรงปกติ 390 (HQ จ่าย 344 บางวัน: ส.ค./ก.ย. 4 วัน)"),
}
REASONS = {"SB-KK030": "หัวหน้ากะ"}
REVIEW = {"SB-KK011": "อ๋อง: อัตรา 350 — ให้ตรวจสอบ"}
TRIAL_RATE = 357.0                                  # the trial-period wage seen on the memos (357 -> 400 after probation)
HQ_PAID = ["SB-FR005", "SB-FR007"]                  # paid by HQ: not part of our payroll
POSITION_FIXES = {"SB-FR071": "เด็กรถ"}


def read_rates(path: str, sheet: str) -> dict:
    import openpyxl
    ws = openpyxl.load_workbook(path)[sheet]
    days_col = next(c.column - 1 for c in ws[4] if c.value == "วัน")
    code_col = next(i for i, c in enumerate(ws[5]) if str(c.value or "").startswith("SB-"))
    out = {}
    for r in range(5, ws.max_row + 1):
        row = [c.value for c in ws[r]]
        code = str(row[code_col] or "")
        if code.startswith("SB-"):
            m = RATE_RE.match(str(row[days_col + 1] or ""))
            out[code.upper()] = float(m.group(1)) if m else None
    return out


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("files", nargs=3, help="drivers, staff, Kaeng Khoi workbooks (in this order)")
    ap.add_argument("--from", dest="start", default="2026-08-26")
    ap.add_argument("--apply", action="store_true")
    args = ap.parse_args()
    found = {}
    for path, sheet in zip(args.files, ("ตารางรายรับ-จ่ายพนง.", "ตารางรายรับ-จ่ายพนง.", "สรุป")):
        found.update(read_rates(path, sheet))
    emps = {e["code"].upper(): e for e in supabase.select_all("employees", {"select": "*", "order": "code.asc"})}
    have = {(r["employee_id"]) for r in supabase.select_all("employee_rates", {"select": "employee_id,effective_from", "effective_from": f"eq.{args.start}", "order": "id.asc"})}
    todo, skipped, notes = [], [], []
    for code, e in emps.items():
        if e["status"] == "resigned" or code in HQ_PAID:
            continue
        rate = OVERRIDES[code][0] if code in OVERRIDES else found.get(code)
        if not rate:
            skipped.append(f"{code} {e['full_name']}: no rate formula in the pay sheets" if code in found else f"{code} {e['full_name']}: not in the pay sheets")
            continue
        if e["id"] in have:
            continue
        review = code in REVIEW or rate == TRIAL_RATE
        note = REVIEW.get(code) or (f"อัตรา {rate:g} (อัตราทดลองงาน?) — ตรวจวันครบทดลองงาน" if rate == TRIAL_RATE else None)
        todo.append((e, {"daily_rate": rate, "effective_from": args.start, "kind": "normal", "reason": REASONS.get(code) or (OVERRIDES[code][1] if code in OVERRIDES else "ตามใบเงินเดือน ก.ย. 2569"),
                         "review": review, "review_note": note}))
    for e, d in todo:
        flag = "  <- REVIEW: " + d["review_note"] if d["review"] else ""
        if d["daily_rate"] != 390 or d["review"]:
            print(f"{e['code']:9} {e['full_name'][:28]:30} {e.get('position') or '':22} {d['daily_rate']:>7g}{flag}")
    print(f"\nrates to write: {len(todo)} (the usual ones are not listed) · skipped: {len(skipped)}")
    for s in skipped:
        print("  skipped:", s)
    print("HQ-paid:", [c for c in HQ_PAID if c in emps], " · position fixes:", POSITION_FIXES)
    if not args.apply:
        print("dry run -- nothing written (add --apply)")
        return
    for e, d in todo:
        supabase.insert("employee_rates", [{"employee_id": e["id"], **rates.clean_rate(d, today=employees.today_th())}])
    for code in HQ_PAID:
        if code in emps:
            supabase.update("employees", {"id": f"eq.{emps[code]['id']}"}, {"pay_group": "hq"})
    for code, pos in POSITION_FIXES.items():
        if code in emps:
            supabase.update("employees", {"id": f"eq.{emps[code]['id']}"}, {"position": pos})
    print(f"written: {len(todo)} rates")


if __name__ == "__main__":
    main()
